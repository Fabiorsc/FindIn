import os
import re
import time
import warnings
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed

import docx
import openpyxl
from pypdf import PdfReader

# Silencia avisos do openpyxl
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

def _ler_texto_plano(caminho):
    linhas = []
    with open(caminho, "r", encoding="utf-8", errors="ignore") as f:
        for num, linha in enumerate(f, start=1):
            if linha.strip():
                linhas.append((num, linha.strip()))
    return linhas

def _ler_docx(caminho):
    linhas = []
    doc = docx.Document(caminho)
    for num, p in enumerate(doc.paragraphs, start=1):
        if p.text.strip():
            linhas.append((num, p.text.strip()))
    return linhas

def _ler_xlsx(caminho):
    linhas = []
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    num_linha = 1
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            conteudo_celulas = [str(celula).strip() for celula in row if celula is not None and str(celula).strip()]
            if conteudo_celulas:
                texto_linha = f"[{sheet_name}] " + " | ".join(conteudo_celulas)
                linhas.append((num_linha, texto_linha))
                num_linha += 1
    wb.close()
    return linhas

def _ler_pdf(caminho):
    linhas = []
    reader = PdfReader(caminho)
    for num_pagina, pagina in enumerate(reader.pages, start=1):
        texto = pagina.extract_text()
        if texto:
            for linha in texto.split("\n"):
                if linha.strip():
                    linhas.append((num_pagina, f"[Pág. {num_pagina}] {linha.strip()}"))
    return linhas

def _buscar_em_arquivo(caminho_arquivo, termo_busca):
    resultados = []
    ext = caminho_arquivo.suffix.lower()

    try:
        if caminho_arquivo.stat().st_size > 30 * 1024 * 1024:  # > 30MB
            return resultados

        if ext == ".docx":
            linhas = _ler_docx(caminho_arquivo)
        elif ext == ".xlsx":
            linhas = _ler_xlsx(caminho_arquivo)
        elif ext == ".pdf":
            linhas = _ler_pdf(caminho_arquivo)
        else:
            linhas = _ler_texto_plano(caminho_arquivo)

        for num_linha, linha in linhas:
            if re.search(re.escape(termo_busca), linha, re.IGNORECASE):
                resultados.append({
                    "arquivo": str(caminho_arquivo),
                    "linha": num_linha,
                    "conteudo": linha
                })

    except Exception:
        pass

    return resultados

def buscar_texto_paralelo(caminho_alvo, termo_busca, extensao="*", callback_progresso=None, cancel_event=None):
    caminho = Path(caminho_alvo)
    if not caminho.exists():
        return [], False

    # Se for um arquivo individual:
    if caminho.is_file():
        if callback_progresso:
            callback_progresso(1, 1, 0)
        return _buscar_em_arquivo(caminho, termo_busca), False

    # Se for um diretório:
    padrao = f"*.{extensao}" if extensao != "*" else "*"
    arquivos = [
        p for p in caminho.rglob(padrao) 
        if p.is_file() and not p.name.startswith("~$")
    ]

    total_arquivos = len(arquivos)
    todos_resultados = []
    processados = 0
    inicio_tempo = time.time()
    foi_cancelado = False

    if total_arquivos == 0:
        return [], False

    with ThreadPoolExecutor(max_workers=os.cpu_count() * 2) as executor:
        futures = {executor.submit(_buscar_em_arquivo, arq, termo_busca): arq for arq in arquivos}
        
        for future in as_completed(futures):
            # Verifica se o utilizador pediu para cancelar
            if cancel_event and cancel_event.is_set():
                foi_cancelado = True
                executor.shutdown(wait=False, cancel_futures=True)
                break

            res = future.result()
            if res:
                todos_resultados.extend(res)
            
            processados += 1
            
            if callback_progresso:
                tempo_decorrido = time.time() - inicio_tempo
                vel = processados / tempo_decorrido if tempo_decorrido > 0 else 0
                restantes = total_arquivos - processados
                eta_segundos = restantes / vel if vel > 0 else 0
                
                callback_progresso(processados, total_arquivos, eta_segundos)

    return todos_resultados, foi_cancelado